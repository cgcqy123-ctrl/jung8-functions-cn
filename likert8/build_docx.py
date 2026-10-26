# -*- coding: utf-8 -*-
"""生成量表版手册（.docx）：likert8/manual/*.md + data/types_*.json + likert8/items.json

用法：python likert8/build_docx.py
输出：荣格八维测验-152题手册.docx（仓库根目录）
"""
import io
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from version import VERSION, DATE  # noqa: E402  唯一的版本号来源

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
HERE = os.path.join(BASE, "likert8")
MANUAL = os.path.join(HERE, "manual")
OUT = os.path.join(BASE, "荣格八维测验-152题手册.docx")

SERIF = "宋体"
SANS = "微软雅黑"
INK = RGBColor(0x1D, 0x20, 0x24)
GRAY = RGBColor(0x6D, 0x71, 0x78)
BODY = 10.5
FUNCS = ["Se", "Si", "Ne", "Ni", "Te", "Ti", "Fe", "Fi"]
FN_CN = {"Se": "外倾感觉", "Si": "内倾感觉", "Ne": "外倾直觉", "Ni": "内倾直觉",
         "Te": "外倾思维", "Ti": "内倾思维", "Fe": "外倾情感", "Fi": "内倾情感"}


# ---------------------------------------------------------------- 基础样式
def set_font(run, cn=SERIF, size=BODY, bold=False, color=None, italic=False):
    run.font.name = cn
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:eastAsia", "w:ascii", "w:hAnsi"):
        fonts.set(qn(attr), cn)


def para(doc, text="", *, indent=True, size=BODY, bold=False, color=None,
         before=0, after=6, align=None, cn=SERIF, line=1.55):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.first_line_indent = Pt(0) if not indent else Cm(0.74)
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align
    add_rich(p, text, cn=cn, size=size, bold=bold, color=color)
    return p


def add_rich(p, text, cn=SERIF, size=BODY, bold=False, color=None):
    """支持 **粗体** 的极简行内标记。"""
    for i, seg in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not seg:
            continue
        set_font(p.add_run(seg), cn=cn, size=size, bold=(bold or i % 2 == 1), color=color)


def h(doc, text, level=1):
    if level == 1:
        doc.add_page_break()
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(0 if level == 1 else 14), Pt(10 if level == 1 else 6)
    pf.keep_with_next = True
    pf.line_spacing = 1.35
    size = {1: 17, 2: 13.5, 3: 11.5}[level]
    add_rich(p, text, cn=SANS, size=size, bold=True)
    if level == 1:
        pb = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "10")
        bottom.set(qn("w:color"), "2F5D50")
        pb.append(bottom)
        p._p.get_or_add_pPr().append(pb)
    return p


def bullet(doc, text, size=BODY):
    p = doc.add_paragraph(style="List Bullet")
    pf = p.paragraph_format
    pf.left_indent = Cm(0.9)
    pf.space_before, pf.space_after = Pt(0), Pt(3)
    pf.line_spacing = 1.5
    add_rich(p, text, size=size)
    return p


def add_table(doc, rows, widths=None, size=9.5, header=True, center_cols=None):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        for ci, cell in enumerate(row):
            c = t.cell(ri, ci)
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.25
            if center_cols and ci in center_cols:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_rich(p, str(cell), size=size, bold=(header and ri == 0))
            if header and ri == 0:
                sh = OxmlElement("w:shd")
                sh.set(qn("w:fill"), "EEEAE0")
                c._tc.get_or_add_tcPr().append(sh)
    if widths:
        for ci, w in enumerate(widths):
            for row in t.rows:
                row.cells[ci].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def footnote_line(doc, text):
    para(doc, text, indent=False, size=9, color=GRAY, after=10, cn=SANS)


def page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    set_font(run, cn=SANS, size=9, color=GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._element.append(fld)


# ---------------------------------------------------------------- md 渲染
def render_md(doc, text, *, first_h1_no_break=True):
    """渲染极简 markdown：H1/H2/H3、段落、列表、表格、<!--page-->。"""
    lines = text.split("\n")
    i, n = 0, len(lines)
    first_h1 = True
    while i < n:
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.strip() == "<!--page-->":
            doc.add_page_break()
            i += 1
            continue
        if line.startswith("### "):
            h(doc, line[4:], 3)
            i += 1
            continue
        if line.startswith("## "):
            h(doc, line[3:], 2)
            i += 1
            continue
        if line.startswith("# "):
            title = line[2:]
            if first_h1 and first_h1_no_break:
                p = doc.add_paragraph()
                pf = p.paragraph_format
                pf.space_before, pf.space_after = Pt(0), Pt(10)
                pf.keep_with_next = True
                add_rich(p, title, cn=SANS, size=17, bold=True)
                pb = OxmlElement("w:pBdr")
                bottom = OxmlElement("w:bottom")
                bottom.set(qn("w:val"), "single")
                bottom.set(qn("w:sz"), "10")
                bottom.set(qn("w:color"), "2F5D50")
                pb.append(bottom)
                p._p.get_or_add_pPr().append(pb)
                first_h1 = False
            else:
                h(doc, title, 1)
            i += 1
            continue
        if line.startswith("|"):
            block = []
            while i < n and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            rows = []
            for b in block:
                cells = [c.strip() for c in b.strip().strip("|").split("|")]
                if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    continue
                rows.append(cells)
            if rows:
                ncol = len(rows[0])
                widths = None
                if ncol == 2:
                    widths = [5.2, 10.5]
                elif ncol == 3:
                    widths = [4.6, 5.0, 6.1]
                elif ncol == 4:
                    widths = [4.2, 4.2, 3.4, 3.9]
                elif ncol == 5:
                    widths = [4.0, 3.4, 3.0, 3.4, 2.7]
                add_table(doc, rows, widths=widths, center_cols=set(range(1, ncol)) if ncol == 2 else None)
            continue
        if line.startswith("- "):
            while i < n and lines[i].startswith("- "):
                bullet(doc, lines[i][2:])
                i += 1
            continue
        buf = []
        while i < n and lines[i].strip() and not lines[i].startswith(("#", "|", "- ", "<!--")):
            buf.append(lines[i].strip())
            i += 1
        para(doc, "".join(buf))


# ---------------------------------------------------------------- 数据
def load_json(*parts):
    with io.open(os.path.join(BASE, *parts), encoding="utf-8") as f:
        return json.load(f)


def load_md(name):
    with io.open(os.path.join(MANUAL, name), encoding="utf-8") as f:
        return f.read()


def build_cover(doc):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    set_font(p.add_run("荣格八维人格测验"), cn=SANS, size=30, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(20)
    set_font(p.add_run("量 表 版"), cn=SANS, size=20, bold=True, color=RGBColor(0x2F, 0x5D, 0x50))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run("中国情境版 · 152 题 · 7 点同意度 · 八维判定"), cn=SANS, size=13, color=GRAY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(60)
    set_font(p.add_run("使用手册"), cn=SANS, size=13, color=GRAY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(VERSION + "　·　" + DATE), cn=SANS, size=11, color=GRAY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    set_font(p.add_run("自我认识工具，不是医学诊断"), cn=SANS, size=10.5, color=GRAY)


def build_toc(doc):
    h(doc, "目录", 1)
    entries = [
        ("前言", 0), ("第一章　为什么还要做一版", 0), ("第二章　荣格的八种功能态度", 0),
        ("第三章　这套题是怎么来的", 0), ("第四章　怎么答", 0), ("第五章　怎么算", 0),
        ("第六章　八种荣格类型", 0), ("第七章　十六型画像", 0), ("第八章　怎么看这个结果", 0),
        ("附录 A　152 题清单", 0), ("附录 B　手工计分表", 0), ("附录 C　参考文献与边界说明", 0),
    ]
    for name, _ in entries:
        para(doc, name, indent=False, after=5, cn=SANS, size=11)
    footnote_line(doc, "页码见每页页脚。")


def build_items_appendix(doc, items):
    h(doc, "附录 A　152 题清单", 1)
    para(doc, "下面按功能分组列出 152 题，题号沿用数据文件里的编号。标「反」的是反向题："
              "答「强烈同意」在该功能上记负分。")
    footnote_line(doc, "题目顺序与网页版不同——网页版按功能交替排列以打散作答定势；"
                       "纸笔作答可以按功能分块做，不影响计分。")
    for fn in FUNCS:
        group = [x for x in items if x["fn"] == fn]
        h(doc, f"{fn}　{FN_CN[fn]}（19 题）", 2)
        for it in group:
            mark = "（反）" if it["reverse"] else ""
            para(doc, f"{it['id']}　{it['text']}{mark}", indent=False, after=3, size=10)


def build_score_sheet(doc):
    h(doc, "附录 B　手工计分表", 1)
    para(doc, "第 1 列填该功能 19 题的分数之和（每题 −3…+3，反向题已翻号）；第 2 列填均值；"
              "第 3 列套公式 (均值 + 3) ÷ 6 × 100；第 4 列填「第 3 列 − 八维均值」。")
    rows = [["功能", "19 题合计", "均值 m", "(m+3)÷6×100", "相对分（减八维均值）"]]
    for fn in FUNCS:
        rows.append([f"{fn}　{FN_CN[fn]}", "", "", "", ""])
    rows.append(["八维均值", "", "", "", ""])
    add_table(doc, rows, widths=[5.4, 3.0, 2.6, 3.4, 4.4], center_cols={1, 2, 3, 4})
    para(doc, "填完之后：相对分最高者＝主导；在「另一类功能 × 相反态度」两个候选里取高者＝辅助；"
              "第三与劣势按第五章 5.6 的对照表推出。", after=4)
    para(doc, "置信度可以粗判：八维里第一名与第二名差 6 分以上，且你觉得自己答得前后一致，可算「高」；"
              "差 3 分以内基本是并列，不要硬分。", after=4)


def build_jung_types(doc, jung):
    h(doc, "第六章　八种荣格类型", 1)
    para(doc, "荣格的类型，是按**主导功能**分的。主导是哪一种功能态度，就属于哪一型——所以八型对应八个功能，"
              "而不是十六型。四字母是在这之上加了一层换算，见第五章 5.7。")
    order = ["ET", "IT", "EF", "IF", "ES", "IS", "EN", "IN"]
    for code in order:
        j = jung.get(code)
        if not j:
            continue
        h(doc, f"{j['name']}（{j['code']}　主导 {j.get('fn', '')}）", 2)
        para(doc, j.get("essence", ""))
        if j.get("everyday"):
            para(doc, "**日常里的样子。**" + j["everyday"])
        if j.get("blind"):
            para(doc, "**盲区与压力。**" + j["blind"])
        if j.get("mistype"):
            para(doc, "**容易和谁混。**" + j["mistype"])


def build_profiles(doc, types):
    h(doc, "第七章　十六型画像", 1)
    para(doc, "下面十六节按四字母排列。每节先给功能栈（这才是判定本身），再给日常三个场景、优势、盲区、压力下的样子、"
              "可以怎么长，以及最容易和哪一型混。")
    order = ["INTJ", "INTP", "ENTJ", "ENTP", "INFJ", "INFP", "ENFJ", "ENFP",
             "ISTJ", "ISFJ", "ESTJ", "ESFJ", "ISTP", "ISFP", "ESTP", "ESFP"]
    for code in order:
        t = types.get(code)
        if not t:
            continue
        stack = "-".join(t.get("stack", []))
        h(doc, f"{code}　{t.get('cn_name', '')}（{stack}）", 2)
        if t.get("core"):
            para(doc, t["core"])
        for key, label in (("work", "职场"), ("family", "家庭与亲戚"), ("love", "亲密关系")):
            if t.get(key):
                para(doc, f"**{label}。**" + t[key])
        if t.get("strength"):
            para(doc, "**优势。**" + "；".join(t["strength"]) + "。")
        if t.get("blindspot"):
            para(doc, "**盲区。**" + "；".join(t["blindspot"]) + "。")
        if t.get("stress"):
            para(doc, "**压力下的样子。**" + t["stress"])
        if t.get("growth"):
            para(doc, "**可以怎么长。**" + t["growth"])
        if t.get("mistype"):
            para(doc, "**容易和谁混。**" + t["mistype"])


# ---------------------------------------------------------------- 主流程
def main() -> int:
    items_doc = load_json("likert8", "items.json")
    items = items_doc["items"]
    if len(items) != 152:
        raise SystemExit(f"题数 {len(items)}，应为 152")

    types, jung = {}, {}
    for name in ("types_a.json", "types_b.json"):
        d = load_json("data", name)
        for t in d["letter_types"]:
            types[t["code"]] = t
        for j in d["jung_types"]:
            jung[j["code"]] = j
    if len(types) != 16 or len(jung) != 8:
        raise SystemExit(f"类型数据不完整：{len(types)} 型 / {len(jung)} 荣格类型")

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2.4), Cm(2.2)
    page_number_footer(sec)

    style = doc.styles["Normal"]
    style.font.name = SERIF
    style.font.size = Pt(BODY)

    build_cover(doc)
    build_toc(doc)

    reading = load_md("03_reading.md")
    head, _, tail = reading.partition("\n# 附录 C")

    render_md(doc, load_md("00_front.md"), first_h1_no_break=False)  # 前言
    render_md(doc, load_md("01_why_theory.md"))     # 第 1–3 章
    render_md(doc, load_md("02_scoring.md"))        # 第 4–5 章
    build_jung_types(doc, jung)                     # 第六章（数据生成）
    build_profiles(doc, types)                      # 第七章（数据生成）
    render_md(doc, head)                            # 第八章
    build_items_appendix(doc, items)                # 附录 A
    build_score_sheet(doc)                          # 附录 B
    render_md(doc, "# 附录 C" + tail)               # 附录 C

    doc.save(OUT)
    print(f"saved: {OUT} ({os.path.getsize(OUT)/1024:.1f} KB)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
