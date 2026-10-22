# -*- coding: utf-8 -*-
"""生成《荣格八维人格测验·中国情境版》Word 文档。

用法：
    python tools/build_docx.py            # 用真实数据构建
    python tools/build_docx.py --demo     # 用合成数据预览排版（不覆盖成品）
"""
from __future__ import annotations

import os
import random
import re
import sys

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from mdparse import parse_md  # noqa: E402
from score import (ATT, CN_NAME, FN_CN, FUNCS, JUNG_OF, OPP, PAIRS,  # noqa: E402
                   STACKS, score, synthetic_items)

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SERIF, SANS = "宋体", "黑体"
BODY = 10.5
GRAY = RGBColor(0x59, 0x59, 0x59)

JUNG_ORDER = ["ET", "IT", "EF", "IF", "ES", "IS", "EN", "IN"]
LETTER_ORDER = ["INTJ", "INTP", "ENTJ", "ENTP", "INFJ", "INFP", "ENFJ", "ENFP",
                "ISTJ", "ISFJ", "ESTJ", "ESFJ", "ISTP", "ISFP", "ESTP", "ESFP"]
CHAPTERS = [
    "写在前面：为什么要重做一版",
    "第一章　为什么现成的网测在你身上总差一点",
    "第二章　理论依据：荣格《心理类型》里的类型是怎么来的",
    "第三章　怎么测",
    "第四章　计分：从 56 个选择到你的类型",
    "第五章　八种荣格类型",
    "第六章　十六型画像",
    "第七章　拿到结果之后，怎么用",
    "第八章　常见问题",
    "附录 A　答题卡",
    "附录 B　计分表",
    "附录 C　参考文献与引文核对",
]


# ---------------------------------------------------------------- 基础排版
def set_font(run, cn=SERIF, en="Times New Roman", size=None, bold=None, color=None):
    run.font.name = en
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), en)
    rfonts.set(qn("w:hAnsi"), en)
    rfonts.set(qn("w:eastAsia"), cn)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_runs(p, text, size=BODY, bold=False, cn=SERIF, color=None):
    import re
    text = text.replace("`", "")
    for part in re.split(r"(\*\*.+?\*\*|\*[^*\n]+?\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = p.add_run(part[2:-2])
            set_font(r, cn="黑体", size=size, bold=True, color=color)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = p.add_run(part[1:-1])
            set_font(r, cn=cn, size=size, bold=bold, color=color)
            r.italic = True
        else:
            r = p.add_run(part)
            set_font(r, cn=cn, size=size, bold=bold, color=color)


def setup_doc() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.3), Cm(2.2)
    sec.left_margin, sec.right_margin = Cm(2.5), Cm(2.5)
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(BODY)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), SERIF)
    pf = st.paragraph_format
    pf.line_spacing = 1.35
    pf.space_before = Pt(0)
    pf.space_after = Pt(4)
    pf.first_line_indent = Pt(BODY * 2)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        s = doc.styles[name]
        s.font.name = "Times New Roman"
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.element.rPr.rFonts.set(qn("w:eastAsia"), SANS)
        s.paragraph_format.first_line_indent = Pt(0)
        s.paragraph_format.keep_with_next = True
    build_footer(sec)
    return doc


def build_footer(section):
    section.different_first_page_header_footer = True
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    r = p.add_run("— ")
    set_font(r, size=9, color=GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE \\* MERGEFORMAT")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), SERIF)
    rpr.append(rfonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "18")
    rpr.append(sz)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = "1"
    run.append(t)
    fld.append(run)
    p._p.append(fld)
    r = p.add_run(" —")
    set_font(r, size=9, color=GRAY)


def h(doc, text, level):
    p = doc.add_heading("", level=level)
    r = p.add_run(text)
    sizes = {1: 17, 2: 13, 3: 11.5}
    set_font(r, cn=SANS, size=sizes[level], bold=True)
    pf = p.paragraph_format
    pf.first_line_indent = Pt(0)
    pf.keep_with_next = True
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.space_before, pf.space_after = Pt(0), Pt(16)
    elif level == 2:
        pf.space_before, pf.space_after = Pt(14), Pt(6)
    else:
        pf.space_before, pf.space_after = Pt(10), Pt(4)
    return p


def para(doc, text="", indent=True, size=BODY, bold=False, align=None,
         before=0, after=4, left=0.0, line=1.35, color=None, keep=False, cn=SERIF):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.first_line_indent = Pt(size * 2) if indent else Pt(0)
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = line
    pf.left_indent = Cm(left)
    pf.keep_with_next = keep
    if align is not None:
        p.alignment = align
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_runs(p, text, size=size, bold=bold, color=color, cn=cn)
    return p


def bullet(doc, text, size=BODY, marker="·"):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.85)
    pf.first_line_indent = Cm(-0.45)
    pf.space_before, pf.space_after = Pt(0), Pt(2)
    pf.line_spacing = 1.3
    r = p.add_run(marker + "　")
    set_font(r, size=size)
    add_runs(p, text, size=size)
    return p


def caption(doc, text):
    p = para(doc, text, indent=False, size=9, align=WD_ALIGN_PARAGRAPH.CENTER,
             before=6, after=2, color=GRAY, keep=True)
    return p


def new_chapter(doc, title):
    doc.add_page_break()
    return h(doc, title, 1)


# ---------------------------------------------------------------- 表格
def shade(cell, fill="F2F2F2"):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def cant_split(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def row_height(row, cm):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:trHeight")
    el.set(qn("w:val"), str(int(cm * 567)))
    el.set(qn("w:hRule"), "atLeast")
    trpr.append(el)


def add_table(doc, rows, widths=None, size=9, header=True, center_cols=None,
              row_h=None):
    ncols = max(len(r) for r in rows)
    if widths is None:
        # 按视觉宽度分配列宽（汉字算 2，ASCII 算 1；可用宽度 16cm，最小 1.4cm）
        def vislen(s):
            s = re.sub(r"[*\s`]", "", s)
            n = sum(2 if ord(c) > 127 else 1 for c in s)
            return n * 1.45 if "http" in s else n
        lens = [0.0] * ncols
        for row in rows:
            for ci in range(ncols):
                cell = row[ci] if ci < len(row) else ""
                lens[ci] = max(lens[ci], min(vislen(cell), 44))
        total = sum(lens) or 1.0
        widths = [max(1.4, 16.0 * v / total) for v in lens]
        scale = 16.0 / sum(widths)
        widths = [w * scale for w in widths]
    t = doc.add_table(rows=len(rows), cols=ncols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tblpr = t._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    grid = t._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(w * 567)))
    for ri, row in enumerate(rows):
        for ci in range(ncols):
            cell = t.cell(ri, ci)
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            txt = row[ci] if ci < len(row) else ""
            add_runs(p, txt, size=size, bold=(header and ri == 0))
            if (header and ri == 0) or (center_cols and ci in center_cols):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if header and ri == 0:
                shade(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths:
                cell.width = Cm(widths[ci])
        if header and ri == 0:
            repeat_header(t.rows[0])
        cant_split(t.rows[ri])
        if row_h:
            row_height(t.rows[ri], row_h)
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(2)
    sp.paragraph_format.line_spacing = 1.0
    return t


# ---------------------------------------------------------------- 渲染 markdown
def render_blocks(doc, blocks, tables=None):
    tables = tables or {}
    for b in blocks:
        t = b["type"]
        if t in ("h1", "h2", "h3"):
            h(doc, b["text"], int(t[1]))
        elif t == "p":
            txt = b["text"].strip()
            if txt.startswith("{{TABLE:") and txt.endswith("}}"):
                key = txt[8:-2].strip()
                if key in tables:
                    tables[key](doc)
                else:
                    para(doc, f"[缺少表格：{key}]")
                continue
            para(doc, txt)
        elif t == "ul":
            for it in b["items"]:
                bullet(doc, it)
        elif t == "ol":
            for i, it in enumerate(b["items"], 1):
                bullet(doc, f"{i}. {it}", marker="　")
        elif t == "quote":
            p = para(doc, b["text"], indent=False, size=9.5, left=0.7, color=GRAY)
            p.paragraph_format.left_indent = Cm(0.7)
        elif t == "table":
            add_table(doc, b["rows"])


SPEC_TO_READER = [
    ("（供 `00_critique.md`、`01_theory.md` 使用，也供成品“附录 C 参考文献”取用。）", ""),
    ("用来支撑 SPEC 的哪一点", "用在哪"),
    ("SPEC §1.1–§1.5", "第二章 2.1–2.5"),
    ("SPEC §1.3–§1.5", "第二章 2.3–2.5"),
    ("SPEC §1.2–§1.5", "第二章 2.2–2.5"),
    ("SPEC §1.4、§1.1", "第二章 2.2、2.4"),
    ("SPEC 未采用八功能模型，见 §1.6", "本书未采用八功能模型（第二章 2.4）"),
    ("SPEC §1 的理论顺序", "本书的理论顺序"),
    ("SPEC §1 边界声明", "第二章 2.6"),
    ("SPEC §1.7", "第二章 2.5、第四章 4.5"),
    ("、§1.8", "、第四章 4.4"),
    ("SPEC §1.8", "第四章 4.4"),
    ("SPEC §1", "第二章"),
    ("`00_critique.md`", "第一章"),
    ("`01_theory.md`", "第二章"),
    ("00 §", "第一章 "),
    ("01 §", "第二章 "),
    ("00 引子", "第一章 引子"),
]


def readerize(text):
    """把研究稿里的内部编号（SPEC、文件名）换成读者看得懂的章节号。"""
    for a, b in SPEC_TO_READER:
        text = text.replace(a, b)
    text = text.replace("§", "")
    for bad in ("SPEC", "§"):
        if bad in text:
            print(f"[readerize] 仍有未替换的内部编号：{bad}")
    return text


def render_md(doc, path, tables=None):
    with open(path, encoding="utf-8") as f:
        blocks = parse_md(f.read())
    if blocks and blocks[0]["type"] == "h1":
        blocks = blocks[1:]
    render_blocks(doc, blocks, tables)


# ---------------------------------------------------------------- 内容：封面与目录
def build_cover(doc, n_items=56):
    for _ in range(4):
        para(doc, "", indent=False, after=0)
    para(doc, "荣格八维人格测验", indent=False, size=30, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, cn=SANS, after=10)
    para(doc, "中 国 情 境 版", indent=False, size=18, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, cn=SANS, after=20)
    p = para(doc, "——基于荣格《心理类型》（1921）的八维认知功能测评",
             indent=False, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, after=26)
    for _ in range(3):
        para(doc, "", indent=False, after=0)
    para(doc, f"{n_items} 个中国生活场景　·　8 种认知功能　·　按荣格的方法定类型",
         indent=False, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    para(doc, "不是四个开关相加，是先定主导、再定辅助",
         indent=False, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, after=60)
    for _ in range(6):
        para(doc, "", indent=False, after=0)
    para(doc, "2026 年 10 月　第 1.1 版", indent=False, size=10,
         align=WD_ALIGN_PARAGRAPH.CENTER, after=6, color=GRAY)
    para(doc, "本测验是自我认识工具，不是医学诊断，也不应作为招聘、筛选或评价他人的依据。",
         indent=False, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, color=GRAY)


def build_toc(doc):
    new_chapter(doc, "目　录")
    for c in CHAPTERS:
        p = para(doc, c, indent=False, size=11, after=6, left=0.6)
        p.paragraph_format.line_spacing = 1.6
    para(doc, "（页码见页脚；电子版可用 Word 的导航窗格或大纲视图快速跳转。）",
         indent=False, size=9, before=14, color=GRAY)


# ---------------------------------------------------------------- 表格数据
def aux_candidates(dom):
    other = [f for f in FUNCS if (f in ("Se", "Si", "Ne", "Ni")) != (dom in ("Se", "Si", "Ne", "Ni"))]
    return [f for f in other if ATT[f] != ATT[dom]]


def code_of(dom, aux):
    if dom in ("Se", "Si", "Ne", "Ni"):
        sn = "N" if dom in ("Ne", "Ni") else "S"
        tf = "T" if aux in ("Te", "Ti") else "F"
    else:
        tf = "T" if dom in ("Te", "Ti") else "F"
        sn = "N" if aux in ("Ne", "Ni") else "S"
    ext = dom if ATT[dom] == "E" else aux
    jp = "J" if ext in ("Te", "Ti", "Fe", "Fi") else "P"
    return f"{ATT[dom]}{sn}{tf}{jp}"


def table_4_1(doc, items):
    caption(doc, "表 4-1　八维计分表（每个功能出现在哪些题、在哪一边）")
    rows = [["功能态度", "出现的题目（题号·该功能在 A 还是 B）", "次数"]]
    for f in FUNCS:
        pos = [f"{it['id']}·{'A' if it['options']['A']['fn'] == f else 'B'}"
               for it in items
               if f in (it["options"]["A"]["fn"], it["options"]["B"]["fn"])]
        rows.append([f"**{f}**　{FN_CN[f]}", "、".join(pos), str(len(pos))])
    add_table(doc, rows, widths=[3.3, 11.3, 1.4], size=8.5, center_cols=[2])


def table_4_2(doc):
    caption(doc, "表 4-2　主导功能决定荣格类型，辅助功能只有两个候选")
    rows = [["主导功能", "荣格类型", "辅助候选（另一类 × 相反态度）", "辅助取前者", "辅助取后者"]]
    for d in ["Te", "Ti", "Fe", "Fi", "Se", "Si", "Ne", "Ni"]:
        cands = aux_candidates(d)
        rows.append([f"**{d}**", JUNG_OF[d][1], "／".join(cands),
                     code_of(d, cands[0]), code_of(d, cands[1])])
    add_table(doc, rows, widths=[2.4, 3.6, 4.6, 2.7, 2.7], size=9, center_cols=[3, 4])


def table_4_3(doc):
    caption(doc, "表 4-3　四字母的换算步骤（以 INTJ 为例）")
    rows = [
        ["步骤", "看什么", "怎么定", "以 INTJ 为例"],
        ["1. E / I", "主导功能的态度", "外倾取 E，内倾取 I", "Ni 是内倾 → I"],
        ["2. S / N", "主导与辅助中属于知觉的那个", "是直觉取 N，是感觉取 S", "主导 Ni → N"],
        ["3. T / F", "主导与辅助中属于判断的那个", "是思维取 T，是情感取 F", "辅助 Te → T"],
        ["4. J / P", "对外展示的功能（主导外倾取主导，否则取辅助）", "判断类取 J，知觉类取 P", "辅助 Te 属判断 → J"],
    ]
    add_table(doc, rows, widths=[2.2, 5.2, 5.0, 3.6], size=9)


def make_example(items, target="INTJ", total=56):
    w = {f: 0.0 for f in FUNCS}
    for wt, f in zip((4.0, 3.0, 2.0, 1.0), STACKS[target]):
        w[f] = wt
    for seed in range(1, 500):
        rnd = random.Random(seed * 7919)
        answers = {}
        for it in items[:total]:
            a, b = it["options"]["A"]["fn"], it["options"]["B"]["fn"]
            d = w[a] - w[b]
            if d > 0:
                base = 2 if d >= 2 else 1
            elif d < 0:
                base = -2 if d <= -2 else -1
            else:
                base = rnd.choice([1, -1])
            if rnd.random() < 0.14:
                base = {2: 1, 1: -1, -1: 1, -2: -1}[base] * (1 if rnd.random() < 0.5 else 1)
                base = rnd.choice([2, 1, -1, -2]) if rnd.random() < 0.4 else base
            answers[it["id"]] = base
        res = score(answers, items[:total])
        if res["code"] == target and res["confidence"] != "低":
            return answers, res
    raise RuntimeError("找不到合适的示例作答")


def table_4_4(doc, items):
    answers, res = make_example(items)
    caption(doc, "表 4-4　示例作答的八维分数（虚构，仅演示算法）")
    rows = [["功能态度", "总分 S（−28 ~ +28）", "展示分 P（0 ~ 100）", "在本人八维中的位置"]]
    order = sorted(FUNCS, key=lambda f: -res["S"][f])
    for f in order:
        rows.append([f"**{f}**　{FN_CN[f]}", f"{res['S'][f]:+.0f}", f"{res['P'][f]:.0f}", str(order.index(f) + 1)])
    add_table(doc, rows, widths=[4.2, 4.0, 4.0, 3.8], size=9, center_cols=[1, 2, 3])
    para(doc, f"这份示例作答推出来的结果是：主导 **{res['dominant']}**，辅助 **{res['auxiliary']}**，"
              f"第三 {res['tertiary']}，劣势 {res['inferior']}，四字母 **{res['code']}**（{res['cn_name']}）。",
         size=9.5, after=6)


def table_6_1(doc):
    caption(doc, "表 6-1　十六型的四个功能位置（按荣格体系，第三功能为迈尔斯补充）")
    rows = [["四字母", "中文名", "主导", "辅助", "第三", "劣势"]]
    for t in LETTER_ORDER:
        st = STACKS[t]
        rows.append([f"**{t}**", CN_NAME[t], st[0], st[1], st[2], st[3]])
    add_table(doc, rows, widths=[2.2, 4.4, 2.3, 2.3, 2.3, 2.3], size=9,
              center_cols=[0, 2, 3, 4, 5])


def table_b_3(doc):
    caption(doc, "表 B-3　主导 + 辅助 → 四字母与功能栈（查表用）")
    rows = [["主导", "辅助", "四字母", "功能栈"]]
    for d in ["Te", "Ti", "Fe", "Fi", "Se", "Si", "Ne", "Ni"]:
        for a in aux_candidates(d):
            code = code_of(d, a)
            rows.append([d, a, f"**{code}**", "-".join(STACKS[code])])
    add_table(doc, rows, widths=[2.4, 2.4, 3.0, 8.2], size=9, center_cols=[0, 1, 2])


# ---------------------------------------------------------------- 内容：题目
def render_items(doc, items):
    parts = ["第一部分（第 1—14 题）", "第二部分（第 15—28 题）",
             "第三部分（第 29—42 题）", "第四部分（第 43—56 题）"]
    for pi in range(4):
        chunk = list(enumerate(items))[pi * 14:(pi + 1) * 14]
        if not chunk:
            continue
        h(doc, parts[pi], 2)
        para(doc, "答题方式：1 = 强烈偏向 A，2 = 略偏向 A，3 = 略微偏向 B，4 = 强烈偏向 B。"
                  "选更接近你**实际反应**的那个，不要选“应该”的那个。",
             indent=False, size=9, after=8, color=GRAY)
        for gi, it in chunk:
            num = gi + 1
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.first_line_indent = Pt(0)
            pf.space_before, pf.space_after = Pt(9), Pt(2)
            pf.line_spacing = 1.3
            pf.keep_with_next = True
            pf.keep_together = True
            r = p.add_run(f"{num:02d}　")
            set_font(r, cn=SANS, size=BODY, bold=True)
            r = p.add_run(it["scenario"])
            set_font(r, cn=SERIF, size=BODY)
            for key in ("A", "B"):
                po = doc.add_paragraph()
                pf = po.paragraph_format
                pf.left_indent = Cm(0.95)
                pf.first_line_indent = Cm(-0.55)
                pf.space_before, pf.space_after = Pt(1), Pt(1)
                pf.line_spacing = 1.25
                pf.keep_together = True
                pf.keep_with_next = (key == "A")
                r = po.add_run(f"{key}　")
                set_font(r, cn=SANS, size=BODY, bold=True)
                r = po.add_run(it["options"][key]["text"])
                set_font(r, cn=SERIF, size=BODY)
        if pi == 1:
            para(doc, "（到这里是第 28 题。可以停一下；回来时从第 29 题接着做，不要回头改前面的答案。）",
                 indent=False, size=9, before=12, color=GRAY)
    para(doc, "做完了？先翻到附录 A 检查答题卡有没有漏填，再进第四章。",
         indent=False, size=9.5, before=16)


# ---------------------------------------------------------------- 内容：类型
def runin(doc, label, text, size=BODY):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.first_line_indent = Pt(size * 2)
    pf.space_before, pf.space_after = Pt(3), Pt(3)
    pf.line_spacing = 1.35
    r = p.add_run(label + "　")
    set_font(r, cn=SANS, size=size, bold=True)
    add_runs(p, text, size=size)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def render_jung_type(doc, jt):
    h(doc, f"{jt['name']}（{jt['code']}）", 2)
    cands = aux_candidates(jt["fn"])
    para(doc, f"**主导** {jt['fn']}（{FN_CN[jt['fn']]}）　·　"
              f"**辅助候选** {'／'.join(cands)}　·　**劣势** {OPP[jt['fn']]}",
         indent=False, size=9, after=4, color=GRAY)
    runin(doc, "本质", jt.get("essence", ""))
    runin(doc, f"辅助支线一（{cands[0]}）", jt.get("aux_a", ""))
    runin(doc, f"辅助支线二（{cands[1]}）", jt.get("aux_b", ""))
    runin(doc, "日常里的样子", jt.get("everyday", ""))
    runin(doc, "盲区与压力", jt.get("blind", ""))
    runin(doc, "别和这些混", jt.get("mistype", ""))


def render_letter_type(doc, t):
    h(doc, f"{t['code']}｜{t['cn_name']}", 2)
    st = t["stack"]
    para(doc, f"功能栈：{st[0]} 主导 · {st[1]} 辅助 · {st[2]} 第三 · {st[3]} 劣势　|　"
              f"荣格归属：{t.get('jung_type', '')}",
         indent=False, size=9, after=2, color=GRAY)
    para(doc, t.get("oneliner", ""), indent=False, size=10.5, bold=True, after=6)
    runin(doc, "核心运作方式", t.get("core", ""))
    runin(doc, "职场", t.get("work", ""))
    runin(doc, "家庭与亲戚", t.get("family", ""))
    runin(doc, "亲密关系", t.get("love", ""))
    para(doc, "优势", indent=False, size=9.5, bold=True, before=5, after=1, cn=SANS)
    for s in t.get("strength", []):
        bullet(doc, s, size=9.5)
    para(doc, "盲区", indent=False, size=9.5, bold=True, before=5, after=1, cn=SANS)
    for s in t.get("blindspot", []):
        bullet(doc, s, size=9.5)
    runin(doc, "压力下的样子", t.get("stress", ""))
    runin(doc, "可以怎么长", t.get("growth", ""))
    runin(doc, "容易和谁混", t.get("mistype", ""))


# ---------------------------------------------------------------- 附录
def build_appendix_a(doc, items):
    new_chapter(doc, "附录 A　答题卡")
    para(doc, "把每题的选择写进对应的格子里：1 = 强烈偏向 A，2 = 略微偏向 A，"
              "3 = 略偏向 B，4 = 强烈偏向 B。做完检查一遍有没有空格。")
    para(doc, "填完之后，翻到附录 B 计分。", after=10)
    caption(doc, "表 A-1　答题卡（完整版 56 题）")
    nums = list(range(1, len(items) + 1))
    rows = [["题号", "选择", "题号", "选择", "题号", "选择", "题号", "选择"]]
    for r in range(14):
        row = []
        for g in range(4):
            idx = r + g * 14
            row += [f"{nums[idx]:02d}" if idx < len(nums) else "", ""]
        rows.append(row)
    add_table(doc, rows, widths=[2.2, 1.8, 2.2, 1.8, 2.2, 1.8, 2.2, 1.8], size=10,
              center_cols=[0, 1, 2, 3, 4, 5, 6, 7], row_h=0.85)
    para(doc, "快速版只做编号为奇数的题（1、3、5……55），共 28 题，也即每个功能对的两个方向里取前一个"
              "（Q01a、Q02a……Q28a）。", indent=False, size=9, color=GRAY)


def build_appendix_b(doc, items):
    new_chapter(doc, "附录 B　计分表")
    para(doc, "纸笔版的计分可以完全手动完成，三步：")
    para(doc, "**第一步**，按表 4-1 把 56 个选择分到八个功能上，各自累加出 S（−28 到 +28）。"
              "下表可以直接用来记：")
    caption(doc, "表 B-1　八维计分汇总（自行填写）")
    rows = [["功能态度", "本题数", "S 小计", "P = 50 + 25 × S ÷ 14", "八维排名"]]
    for f in FUNCS:
        rows.append([f"**{f}**　{FN_CN[f]}", "14", "", "", ""])
    add_table(doc, rows, widths=[4.0, 2.4, 3.0, 4.4, 2.2], size=9,
              center_cols=[1, 2, 3, 4], row_h=0.8)
    para(doc, "**第二步**，S 最高的那个功能就是你的主导功能。按表 4-2 找到它的荣格类型和两个辅助候选，"
              "再比较这两个候选的 S，高者即辅助。")
    para(doc, "**第三步**，用下表把主导 + 辅助换成四字母，再对照第六章找到你的画像。")
    para(doc, "注意：八维分数是**相对强度**。八个 S 相加恒等于 0，所以 P 的绝对值没有意义，"
              "只有排序有意义。不要拿它和别人的分数比。", size=9.5)
    para(doc, "快速版把表 B-1 的 14 改成 7、把公式里除以 14 改成除以 7 即可，判定流程完全相同。",
         indent=False, size=9, color=GRAY)
    doc.add_page_break()
    table_b_3(doc)


# ---------------------------------------------------------------- 主流程
def load_data(demo=False):
    if demo:
        items = synthetic_items()
        for it in items:
            it["scenario"] = f"（示例场景 {it['id']}）一个具体到可以想象的中国日常情境。"
            it["options"] = {"A": {"fn": it["options"]["A"]["fn"], "text": "（示例选项 A 的文字，约二十余字。）"},
                             "B": {"fn": it["options"]["B"]["fn"], "text": "（示例选项 B 的文字，约二十余字。）"}}
        items.sort(key=lambda x: x["id"])
        letter = []
        for code in LETTER_ORDER:
            st = STACKS[code]
            letter.append({"code": code, "cn_name": CN_NAME[code], "stack": st,
                           "jung_type": JUNG_OF[st[0]][1], "oneliner": "（一句话画像）",
                           "core": "（核心运作方式，示例文字。）", "work": "（职场。）",
                           "family": "（家庭与亲戚。）", "love": "（亲密关系。）",
                           "strength": ["（优势一）", "（优势二）", "（优势三）"],
                           "blindspot": ["（盲区一）", "（盲区二）", "（盲区三）"],
                           "stress": "（压力下的样子。）", "growth": "（可以怎么长。）",
                           "mistype": "（容易和谁混。）"})
        jung = []
        for code, name, fn in [("ET", "外倾思维型", "Te"), ("IT", "内倾思维型", "Ti"),
                               ("EF", "外倾情感型", "Fe"), ("IF", "内倾情感型", "Fi"),
                               ("ES", "外倾感觉型", "Se"), ("IS", "内倾感觉型", "Si"),
                               ("EN", "外倾直觉型", "Ne"), ("IN", "内倾直觉型", "Ni")]:
            jung.append({"code": code, "name": name, "fn": fn, "essence": "（本质。）",
                         "aux_a": "（辅助支线一。）", "aux_b": "（辅助支线二。）",
                         "everyday": "（日常里的样子。）", "blind": "（盲区与压力。）",
                         "mistype": "（别和这些混。）"})
        types = {"letter_types": letter, "jung_types": jung}
    else:
        with open(os.path.join(BASE, "data", "items_a.json"), encoding="utf-8") as f:
            items = __import__("json").load(f)["items"]
        with open(os.path.join(BASE, "data", "items_b.json"), encoding="utf-8") as f:
            items += __import__("json").load(f)["items"]
        items.sort(key=lambda x: x["id"])
        types = {"letter_types": [], "jung_types": []}
        for name in ("types_a.json", "types_b.json"):
            with open(os.path.join(BASE, "data", name), encoding="utf-8") as f:
                d = __import__("json").load(f)
            types["letter_types"] += d.get("letter_types", [])
            types["jung_types"] += d.get("jung_types", [])
    return items, types


def build(demo=False) -> str:
    items, types = load_data(demo)
    doc = setup_doc()
    build_cover(doc, len(items))
    build_toc(doc)

    def md(name, title, transform=None):
        new_chapter(doc, title)
        path = os.path.join(BASE, name)
        if os.path.exists(path):
            with open(path, encoding="utf-8") as f:
                text = f.read()
            if transform:
                text = transform(text)
            blocks = parse_md(text)
            if blocks and blocks[0]["type"] == "h1":
                blocks = blocks[1:]
            render_blocks(doc, blocks, tables)
        else:
            para(doc, f"〔缺少文稿：{name}〕", indent=False)

    tables = {}

    md("data/front.md", CHAPTERS[0])
    md("research/00_critique.md", CHAPTERS[1])
    md("research/01_theory.md", CHAPTERS[2])

    # 第三章
    new_chapter(doc, CHAPTERS[3])
    render_md(doc, os.path.join(BASE, "data", "ch3.md"), tables)
    h(doc, "3.5　题目（完整版 56 题）", 2)
    para(doc, "下面 56 题分为四部分。请一次做完前两部分，休息后再做后两部分；"
              "纸笔版请把答案写在附录 A 的答题卡上。", before=2, after=6)
    render_items(doc, items)

    # 第四章
    tables["table_4_1"] = lambda d: table_4_1(d, items)
    tables["table_4_2"] = table_4_2
    tables["table_4_3"] = table_4_3
    tables["table_4_4"] = lambda d: table_4_4(d, items)
    md("data/ch4.md", CHAPTERS[4])

    # 第五章
    new_chapter(doc, CHAPTERS[5])
    render_md(doc, os.path.join(BASE, "data", "ch5_intro.md"), tables)
    by_code = {j["code"]: j for j in types["jung_types"]}
    for code in JUNG_ORDER:
        if code in by_code:
            render_jung_type(doc, by_code[code])
        else:
            para(doc, f"〔缺少荣格类型：{code}〕", indent=False)

    # 第六章
    tables["table_6_1"] = table_6_1
    new_chapter(doc, CHAPTERS[6])
    render_md(doc, os.path.join(BASE, "data", "ch6_intro.md"), tables)
    by_letter = {t["code"]: t for t in types["letter_types"]}
    for code in LETTER_ORDER:
        if code in by_letter:
            render_letter_type(doc, by_letter[code])
        else:
            para(doc, f"〔缺少类型：{code}〕", indent=False)

    md("data/ch7.md", CHAPTERS[7])
    md("data/ch8.md", CHAPTERS[8])

    build_appendix_a(doc, items)
    build_appendix_b(doc, items)
    md("research/02_sources.md", CHAPTERS[11], transform=readerize)

    para(doc, "— 全 书 完 —", indent=False, size=11, align=WD_ALIGN_PARAGRAPH.CENTER,
         before=48, color=GRAY)

    out = os.path.join(BASE, "PREVIEW_排版预览.docx" if demo else "荣格八维人格测验-中国情境版.docx")
    doc.save(out)
    return out


if __name__ == "__main__":
    demo = "--demo" in sys.argv
    path = build(demo)
    print("saved:", path)
